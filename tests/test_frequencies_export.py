"""
Tests for Frequency Export functionality.

This file tests the frequency and count sheets added to the Excel export:
- Counts_All (wide format, counts only)
- Counts_Selected (wide format, counts only)
- Frequencies_All (wide format, percentages only)
- Frequencies_Selected (wide format, percentages only)
- Optional tidy format sheets for counts and frequencies
"""

import pytest
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import load_workbook
import pandas as pd

from flowkit_gui.models.workspace_model import Sample, GateInfo
from flowkit_gui.models.keywords_model import KeywordsModel
from flowkit_gui.models.statistics_model import MFIMetric, StatisticsModel, FrequencyDataType
from flowkit_gui.exporters.xlsx_exporter import XLSXExporter


# Extended Mock WorkspaceModel with gate statistics support
class MockWorkspaceModel:
    """Mock workspace model for testing frequency exports."""

    def __init__(self):
        self._samples = []
        self._gates = []
        self._channels = []
        self._gate_stats = {}  # {sample_id: {gate_id: {count, percent_parent}}}

    @property
    def samples(self):
        return self._samples

    @property
    def gates(self):
        return self._gates

    @property
    def channels(self):
        return self._channels

    def set_samples(self, samples):
        self._samples = samples

    def set_gates(self, gates):
        self._gates = gates

    def set_channels(self, channels):
        self._channels = channels

    def set_gate_stats(self, sample_id: str, gate_id: str, count: int, percent_parent: float):
        """Set mock gate statistics for testing."""
        if sample_id not in self._gate_stats:
            self._gate_stats[sample_id] = {}
        self._gate_stats[sample_id][gate_id] = {
            'count': count,
            'percent_parent': percent_parent,
            'percent_total': percent_parent  # Simplified for testing
        }

    def get_gate_statistics(self, sample_id: str, gate_id: str):
        """Get gate statistics (mock implementation)."""
        if sample_id in self._gate_stats and gate_id in self._gate_stats[sample_id]:
            return self._gate_stats[sample_id][gate_id]
        return {'count': 0, 'percent_parent': 0.0, 'percent_total': 0.0}

    def has_fcs_data(self, sample_id: str) -> bool:
        """Check if sample has FCS data."""
        for sample in self._samples:
            if sample.sample_id == sample_id:
                return sample.event_count > 0
        return False

    def get_samples_with_data(self):
        """Get samples that have FCS data."""
        return [s for s in self._samples if s.event_count > 0]

    def get_samples_without_data(self):
        """Get samples that are missing FCS data."""
        return [s for s in self._samples if s.event_count == 0]

    def get_missing_gates_info(self, gate_ids):
        """Get info about gates missing from samples (mock returns empty dict)."""
        return {}


def create_test_sample(sample_id, name, keywords):
    """Create a test sample."""
    return Sample(
        sample_id=sample_id,
        name=name,
        fcs_path=None,
        group="",
        keywords=keywords,
        event_count=10000
    )


def create_test_gate(gate_id, name, path, gate_type="PolygonGate"):
    """Create a test gate."""
    return GateInfo(
        gate_id=gate_id,
        name=name,
        path=path,
        gate_type=gate_type,
        parameters=["FSC-A", "SSC-A"]
    )


class TestCountsSheets:
    """Tests for counts sheets (new feature)."""

    def test_counts_all_sheet_created(self, tmp_path):
        """Test that Counts_All sheet is created when include_all_sheets=True."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Counts_All" in wb.sheetnames

    def test_counts_selected_sheet_created_when_markers_selected(self, tmp_path):
        """Test that Counts_Selected sheet is created when markers are selected."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4"],  # Markers selected
            all_markers=["CD4"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Counts_Selected" in wb.sheetnames

    def test_counts_sheet_has_only_count_columns(self, tmp_path):
        """Test that counts sheet only has count columns, not percentages."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)
        workspace.set_gate_stats("s1", "live", 6000, 75.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells", "live"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Counts_All"]

        headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
        assert "Cells_count" in headers
        assert "Cells/Live_count" in headers
        assert "Cells_%parent" not in headers
        assert "Cells/Live_%parent" not in headers


class TestFrequencyWideSheets:
    """Tests for wide format frequency sheets."""

    def test_frequencies_all_sheet_created(self, tmp_path):
        """Test that Frequencies_All sheet is created when include_all_sheets=True."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Frequencies_All" in wb.sheetnames

    def test_frequencies_selected_sheet_created_when_markers_selected(self, tmp_path):
        """Test that Frequencies_Selected sheet is created when markers are selected."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4"],  # Markers selected
            all_markers=["CD4"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Frequencies_Selected" in wb.sheetnames

    def test_frequencies_wide_column_naming(self, tmp_path):
        """Test that wide format uses full path column names."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)
        workspace.set_gate_stats("s1", "live", 6000, 75.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells", "live"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)

        # Check Counts_All has only count columns
        ws_counts = wb["Counts_All"]
        counts_headers = [ws_counts.cell(1, i).value for i in range(1, ws_counts.max_column + 1)]
        assert "Cells_count" in counts_headers
        assert "Cells/Live_count" in counts_headers
        assert "Cells_%parent" not in counts_headers
        assert "Cells/Live_%parent" not in counts_headers

        # Check Frequencies_All has only percent columns
        ws_freq = wb["Frequencies_All"]
        freq_headers = [ws_freq.cell(1, i).value for i in range(1, ws_freq.max_column + 1)]
        assert "Cells_%parent" in freq_headers
        assert "Cells/Live_%parent" in freq_headers
        assert "Cells_count" not in freq_headers
        assert "Cells/Live_count" not in freq_headers

    def test_frequencies_wide_data_values(self, tmp_path):
        """Test that wide format has correct data values in separate sheets."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)

        # Check counts in Counts_All
        ws_counts = wb["Counts_All"]
        counts_headers = [ws_counts.cell(1, i).value for i in range(1, ws_counts.max_column + 1)]
        count_col = counts_headers.index("Cells_count") + 1
        assert ws_counts.cell(2, count_col).value == 5000

        # Check percentages in Frequencies_All
        ws_freq = wb["Frequencies_All"]
        freq_headers = [ws_freq.cell(1, i).value for i in range(1, ws_freq.max_column + 1)]
        pct_col = freq_headers.index("Cells_%parent") + 1
        assert ws_freq.cell(2, pct_col).value == 50.0


class TestFrequencyTidySheets:
    """Tests for tidy format frequency sheets."""

    def test_tidy_sheets_not_included_by_default(self, tmp_path):
        """Test tidy sheets are not included when checkbox is off."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # Counts tidy sheets not included
        assert "Counts_All_Tidy" not in wb.sheetnames
        assert "Counts_Selected_Tidy" not in wb.sheetnames
        # Frequencies tidy sheets not included
        assert "Frequencies_All_Tidy" not in wb.sheetnames
        assert "Frequencies_Selected_Tidy" not in wb.sheetnames

    def test_tidy_sheets_included_when_enabled(self, tmp_path):
        """Test tidy frequency and count sheets are included when enabled with markers."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4"],  # Markers selected for _Selected sheets
            all_markers=["CD4"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=True,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # Counts tidy sheets
        assert "Counts_All_Tidy" in wb.sheetnames
        assert "Counts_Selected_Tidy" in wb.sheetnames
        # Frequencies tidy sheets
        assert "Frequencies_All_Tidy" in wb.sheetnames
        assert "Frequencies_Selected_Tidy" in wb.sheetnames

    def test_tidy_format_structure(self, tmp_path):
        """Test tidy format has correct structure with separate count and frequency sheets."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)
        workspace.set_gate_stats("s1", "live", 6000, 75.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells", "live"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=True,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)

        # Check Counts_All_Tidy headers
        ws_counts = wb["Counts_All_Tidy"]
        counts_headers = [ws_counts.cell(1, i).value for i in range(1, ws_counts.max_column + 1)]
        assert "sample_id" in counts_headers
        assert "sample_name" in counts_headers
        assert "gate_path" in counts_headers
        assert "gate_name" in counts_headers
        assert "count" in counts_headers
        assert "%parent" not in counts_headers

        # Check Frequencies_All_Tidy headers
        ws_freq = wb["Frequencies_All_Tidy"]
        freq_headers = [ws_freq.cell(1, i).value for i in range(1, ws_freq.max_column + 1)]
        assert "sample_id" in freq_headers
        assert "sample_name" in freq_headers
        assert "gate_path" in freq_headers
        assert "gate_name" in freq_headers
        assert "%parent" in freq_headers
        assert "count" not in freq_headers

    def test_tidy_format_data_rows(self, tmp_path):
        """Test tidy format has one row per gate per sample."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"}),
            create_test_sample("s2", "Sample2.fcs", {"Mouse #": "456"})
        ])
        # Set stats for both samples
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)
        workspace.set_gate_stats("s1", "live", 6000, 75.0)
        workspace.set_gate_stats("s2", "cells", 7000, 70.0)
        workspace.set_gate_stats("s2", "live", 5000, 71.4)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells", "live"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=True,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Frequencies_All_Tidy"]

        # Should have header + 4 data rows (2 samples x 2 gates)
        assert ws.max_row == 5


class TestFrequencyEdgeCases:
    """Tests for edge cases in frequency calculations."""

    def test_empty_gate_shows_zero(self, tmp_path):
        """Test that empty gates show count=0 and %parent=0 in respective sheets."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 0, 0.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)

        # Check count in Counts_All
        ws_counts = wb["Counts_All"]
        counts_headers = [ws_counts.cell(1, i).value for i in range(1, ws_counts.max_column + 1)]
        count_col = counts_headers.index("Cells_count") + 1
        assert ws_counts.cell(2, count_col).value == 0

        # Check percent in Frequencies_All
        ws_freq = wb["Frequencies_All"]
        freq_headers = [ws_freq.cell(1, i).value for i in range(1, ws_freq.max_column + 1)]
        pct_col = freq_headers.index("Cells_%parent") + 1
        assert ws_freq.cell(2, pct_col).value == 0.0

    def test_root_gate_percent_parent(self, tmp_path):
        """Test that root gate shows 100% for %parent in Frequencies_All."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        # Root gate should be 100% of parent (itself is the root)
        workspace.set_gate_stats("s1", "cells", 10000, 100.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Frequencies_All"]

        headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
        pct_col = headers.index("Cells_%parent") + 1

        assert ws.cell(2, pct_col).value == 100.0

    def test_no_gates_selected(self, tmp_path):
        """Test export works with no gates selected."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],  # No gates selected
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # _All sheets should exist when include_all_sheets=True
        assert "Frequencies_All" in wb.sheetnames
        # _Selected sheets should NOT exist (no markers selected)
        assert "Frequencies_Selected" not in wb.sheetnames


class TestFrequencySheetOrder:
    """Tests for correct sheet ordering in workbook."""

    def test_sheet_order_without_tidy(self, tmp_path):
        """Test correct sheet order without tidy sheets (with markers selected)."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4"],  # Markers selected for _Selected sheets
            all_markers=["CD4"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        expected_order = [
            "Gating_Strategy",
            "Keywords",
            "Counts_All",
            "Counts_Selected",
            "Frequencies_All",
            "Frequencies_Selected",
            "MFI_Selected",
            "MFI_All",
            "Concatenated_Selected",
            "Concatenated_All",
            "Warnings"  # Added due to MFI calculation on mock
        ]
        assert wb.sheetnames == expected_order

    def test_sheet_order_with_tidy(self, tmp_path):
        """Test correct sheet order with tidy sheets (with markers selected)."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4"],  # Markers selected for _Selected sheets
            all_markers=["CD4"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_tidy_frequencies=True,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        expected_order = [
            "Gating_Strategy",
            "Keywords",
            "Counts_All",
            "Counts_Selected",
            "Frequencies_All",
            "Frequencies_Selected",
            "Counts_All_Tidy",
            "Counts_Selected_Tidy",
            "Frequencies_All_Tidy",
            "Frequencies_Selected_Tidy",
            "MFI_Selected",
            "MFI_All",
            "Concatenated_Selected",
            "Concatenated_All",
            "Warnings"  # Added due to MFI calculation on mock
        ]
        assert wb.sheetnames == expected_order

    def test_sheet_order_with_keywords_wide(self, tmp_path):
        """Test correct sheet order with Keywords_Wide included (with markers selected)."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4"],  # Markers selected for _Selected sheets
            all_markers=["CD4"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=True,
            include_tidy_frequencies=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        expected_order = [
            "Gating_Strategy",
            "Keywords",
            "Keywords_Wide",
            "Counts_All",
            "Counts_Selected",
            "Frequencies_All",
            "Frequencies_Selected",
            "MFI_Selected",
            "MFI_All",
            "Concatenated_Selected",
            "Concatenated_All",
            "Warnings"  # Added due to MFI calculation on mock
        ]
        assert wb.sheetnames == expected_order


class TestStatisticsModelFrequencies:
    """Tests for StatisticsModel frequency calculation methods."""

    def test_calculate_frequencies_wide_both(self):
        """Test wide format frequency calculation with both counts and percentages."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)
        workspace.set_gate_stats("s1", "live", 6000, 75.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells", "live"], wide=True, data_type=FrequencyDataType.BOTH)

        # Should have one row
        assert len(df) == 1

        # Check columns - both count and percent
        assert "sample_id" in df.columns
        assert "sample_name" in df.columns
        assert "Cells_count" in df.columns
        assert "Cells_%parent" in df.columns
        assert "Cells/Live_count" in df.columns
        assert "Cells/Live_%parent" in df.columns

        # Check values
        assert df.iloc[0]["Cells_count"] == 8000
        assert df.iloc[0]["Cells_%parent"] == 80.0
        assert df.iloc[0]["Cells/Live_count"] == 6000
        assert df.iloc[0]["Cells/Live_%parent"] == 75.0

    def test_calculate_frequencies_wide_counts_only(self):
        """Test wide format frequency calculation with counts only."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells"], wide=True, data_type=FrequencyDataType.COUNTS)

        # Check has count but not percent
        assert "Cells_count" in df.columns
        assert "Cells_%parent" not in df.columns
        assert df.iloc[0]["Cells_count"] == 8000

    def test_calculate_frequencies_wide_percent_only(self):
        """Test wide format frequency calculation with percent only."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells"], wide=True, data_type=FrequencyDataType.PERCENT)

        # Check has percent but not count
        assert "Cells_%parent" in df.columns
        assert "Cells_count" not in df.columns
        assert df.iloc[0]["Cells_%parent"] == 80.0

    def test_calculate_frequencies_tidy_both(self):
        """Test tidy format frequency calculation with both counts and percentages."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)
        workspace.set_gate_stats("s1", "live", 6000, 75.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells", "live"], wide=False, data_type=FrequencyDataType.BOTH)

        # Should have two rows (one per gate)
        assert len(df) == 2

        # Check columns - both count and percent
        assert "sample_id" in df.columns
        assert "sample_name" in df.columns
        assert "gate_path" in df.columns
        assert "gate_name" in df.columns
        assert "count" in df.columns
        assert "%parent" in df.columns

        # Check values for first gate
        cells_row = df[df["gate_name"] == "cells"].iloc[0]
        assert cells_row["count"] == 8000
        assert cells_row["%parent"] == 80.0

    def test_calculate_frequencies_tidy_counts_only(self):
        """Test tidy format frequency calculation with counts only."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells"], wide=False, data_type=FrequencyDataType.COUNTS)

        # Check has count but not percent
        assert "count" in df.columns
        assert "%parent" not in df.columns
        assert df.iloc[0]["count"] == 8000

    def test_calculate_frequencies_tidy_percent_only(self):
        """Test tidy format frequency calculation with percent only."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])
        workspace.set_gate_stats("s1", "cells", 8000, 80.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells"], wide=False, data_type=FrequencyDataType.PERCENT)

        # Check has percent but not count
        assert "%parent" in df.columns
        assert "count" not in df.columns
        assert df.iloc[0]["%parent"] == 80.0

    def test_calculate_frequencies_with_keywords(self):
        """Test that keywords are included in output."""
        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123", "Tumor": "A"})
        ])
        workspace.set_gate_stats("s1", "cells", 5000, 50.0)

        stats_model = StatisticsModel(workspace)
        df = stats_model.calculate_frequencies(["cells"], wide=True, include_keywords=True)

        assert "Mouse #" in df.columns
        assert "Tumor" in df.columns
        assert df.iloc[0]["Mouse #"] == "123"
        assert df.iloc[0]["Tumor"] == "A"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
