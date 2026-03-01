"""
Tests for XLSX Exporter.
"""

import pytest
from pathlib import Path
import sys
import tempfile
import os

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import load_workbook
import pandas as pd

from flowkit_gui.models.workspace_model import Sample, GateInfo
from flowkit_gui.models.keywords_model import KeywordsModel
from flowkit_gui.models.statistics_model import MFIMetric
from flowkit_gui.exporters.xlsx_exporter import XLSXExporter


# Mock WorkspaceModel for testing without FlowKit dependency
class MockWorkspaceModel:
    """Mock workspace model for testing."""

    def __init__(self):
        self._samples = []
        self._gates = []
        self._channels = []

    @property
    def samples(self):
        return self._samples

    @property
    def gates(self):
        return self._gates

    @property
    def channels(self):
        return self._channels

    def set_samples(self, samples):
        self._samples = samples

    def set_gates(self, gates):
        self._gates = gates

    def set_channels(self, channels):
        self._channels = channels

    def has_fcs_data(self, sample_id: str) -> bool:
        """Check if sample has FCS data (mock always returns True for samples with event_count > 0)."""
        for sample in self._samples:
            if sample.sample_id == sample_id:
                return sample.event_count > 0
        return False

    def get_samples_with_data(self):
        """Get samples that have FCS data."""
        return [s for s in self._samples if s.event_count > 0]

    def get_samples_without_data(self):
        """Get samples that are missing FCS data."""
        return [s for s in self._samples if s.event_count == 0]

    def get_gate_statistics(self, sample_id: str, gate_id: str):
        """Get gate statistics (mock returns zeros)."""
        return {'count': 0, 'percent_parent': 0.0, 'percent_total': 0.0}

    def get_missing_gates_info(self, gate_ids):
        """Get info about gates missing from samples (mock returns empty dict)."""
        return {}

    def get_workspace_compensation_info(self):
        """Get compensation info (mock returns empty structure)."""
        return {
            'matrices': {},
            'assignments': {}
        }

    def get_sample_by_id(self, sample_id: str):
        """Get a sample by its ID."""
        for sample in self._samples:
            if sample.sample_id == sample_id:
                return sample
        return None


def create_test_sample(sample_id, name, keywords):
    """Create a test sample."""
    return Sample(
        sample_id=sample_id,
        name=name,
        fcs_path=None,
        group="",
        keywords=keywords,
        event_count=10000
    )


def create_test_gate(gate_id, name, path, gate_type="PolygonGate"):
    """Create a test gate."""
    return GateInfo(
        gate_id=gate_id,
        name=name,
        path=path,
        gate_type=gate_type,
        parameters=["FSC-A", "SSC-A"]
    )


class TestXLSXExporterGatingStrategy:
    """Tests for gating strategy sheet generation."""

    def test_gating_strategy_headers(self, tmp_path):
        """Test that gating strategy sheet has correct headers."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Gating_Strategy"]

        # Check headers
        assert ws.cell(1, 1).value == "Gate Path"
        assert ws.cell(1, 2).value == "Gate Name"
        assert ws.cell(1, 3).value == "Depth"
        assert ws.cell(1, 4).value == "Gate Type"
        assert ws.cell(1, 5).value == "Parameters"

    def test_gating_strategy_full_paths(self, tmp_path):
        """Test that gate paths are verbose (not tree characters)."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("single", "Single Cells", ("Cells", "Single Cells")),
            create_test_gate("live", "Live", ("Cells", "Single Cells", "Live")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells", "single", "live"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Gating_Strategy"]

        # Check full paths (row 2, 3, 4 after header)
        assert ws.cell(2, 1).value == "Cells"
        assert ws.cell(3, 1).value == "Cells/Single Cells"
        assert ws.cell(4, 1).value == "Cells/Single Cells/Live"


class TestXLSXExporterKeywords:
    """Tests for keywords sheets."""

    def test_keywords_tidy_sheet_exists(self, tmp_path):
        """Test that Keywords (tidy) sheet is created."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123", "Tumor": "C"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Keywords" in wb.sheetnames

    def test_keywords_tidy_format(self, tmp_path):
        """Test tidy format has correct structure."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123", "Tumor": "C"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Keywords"]

        # Check headers
        assert ws.cell(1, 1).value == "sample_id"
        assert ws.cell(1, 2).value == "sample_name"
        assert ws.cell(1, 3).value == "keyword_name"
        assert ws.cell(1, 4).value == "keyword_value"

    def test_keywords_wide_not_included_by_default(self, tmp_path):
        """Test Keywords_Wide is not included when checkbox is off."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Keywords_Wide" not in wb.sheetnames

    def test_keywords_wide_included_when_enabled(self, tmp_path):
        """Test Keywords_Wide is included when checkbox is on."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123", "Tumor": "C"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "Keywords_Wide" in wb.sheetnames


class TestXLSXExporterStatistics:
    """Tests for statistics sheets."""

    def test_mfi_selected_sheet_exists(self, tmp_path):
        """Test MFI_Selected sheet is created."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=["CD4", "CD8"],
            all_markers=["CD4", "CD8", "CD45"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "MFI_Selected" in wb.sheetnames
        assert "Concatenated_Selected" in wb.sheetnames

    def test_mfi_all_sheet_exists(self, tmp_path):
        """Test MFI_All sheet is created when include_all_sheets=True."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=["CD4", "CD8", "CD45"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        assert "MFI_All" in wb.sheetnames
        assert "Concatenated_All" in wb.sheetnames


class TestXLSXExporterFormatting:
    """Tests for Excel formatting."""

    def test_header_row_frozen(self, tmp_path):
        """Test that header row is frozen."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",))
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        ws = wb["Gating_Strategy"]

        # Check freeze pane is set
        assert ws.freeze_panes == "A2"

    def test_progress_callback(self, tmp_path):
        """Test progress callback is called."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        progress_calls = []

        def progress_callback(current, total):
            progress_calls.append((current, total))

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath), progress_callback=progress_callback)

        # Should have progress calls
        assert len(progress_calls) > 0
        # Final call should be complete
        assert progress_calls[-1][0] == progress_calls[-1][1]


class TestXLSXExporterSheetCount:
    """Tests for correct number of sheets."""

    def test_sheet_count_without_wide(self, tmp_path):
        """Test correct number of sheets without Keywords_Wide but with include_all_sheets."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # Should have: Gating_Strategy, Keywords, Compensation,
        # Counts_All, Frequencies_All, MFI_All, Concatenated_All,
        # Counts_Selected, Frequencies_Selected (count/freq always written)
        # (no MFI_Selected or Concatenated_Selected when selected_markers is empty)
        assert len(wb.sheetnames) == 9

    def test_sheet_count_with_wide(self, tmp_path):
        """Test correct number of sheets with Keywords_Wide and include_all_sheets."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=True,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # Should have: Gating_Strategy, Keywords, Keywords_Wide, Compensation,
        # Counts_All, Frequencies_All, MFI_All, Concatenated_All,
        # Counts_Selected, Frequencies_Selected (count/freq always written)
        # (no MFI_Selected or Concatenated_Selected when selected_markers is empty)
        assert len(wb.sheetnames) == 10

    def test_selected_sheets_included_when_markers_selected(self, tmp_path):
        """Test _Selected sheets ARE included when markers are selected."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",))
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=["CD4", "CD8"],  # Markers ARE selected
            all_markers=["CD4", "CD8", "CD45"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # Should have all 11 sheets including _Selected and _All variants + Compensation
        assert "MFI_Selected" in wb.sheetnames
        assert "Concatenated_Selected" in wb.sheetnames
        assert "Counts_Selected" in wb.sheetnames
        assert "Frequencies_Selected" in wb.sheetnames
        assert "MFI_All" in wb.sheetnames
        assert "Concatenated_All" in wb.sheetnames
        assert "Counts_All" in wb.sheetnames
        assert "Frequencies_All" in wb.sheetnames
        assert len(wb.sheetnames) == 11

    def test_selected_sheets_excluded_when_no_markers(self, tmp_path):
        """Test MFI/Concatenated _Selected sheets are excluded when no markers, but count/freq always present."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",))
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],  # No markers selected
            all_markers=["CD4", "CD8", "CD45"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False,
            include_all_sheets=True
        )

        filepath = tmp_path / "test.xlsx"
        exporter.export(str(filepath))

        wb = load_workbook(filepath)
        # MFI/Concatenated _Selected sheets should NOT exist (need markers)
        assert "MFI_Selected" not in wb.sheetnames
        assert "Concatenated_Selected" not in wb.sheetnames
        # Count/Frequency _Selected sheets SHOULD exist (gate-level, no markers needed)
        assert "Counts_Selected" in wb.sheetnames
        assert "Frequencies_Selected" in wb.sheetnames
        # _All sheets should exist (when include_all_sheets=True) + Compensation
        assert "MFI_All" in wb.sheetnames
        assert "Concatenated_All" in wb.sheetnames
        assert "Counts_All" in wb.sheetnames
        assert "Frequencies_All" in wb.sheetnames
        assert len(wb.sheetnames) == 9


class TestXLSXExporterSummary:
    """Tests for export summary generation."""

    def test_export_returns_summary(self, tmp_path):
        """Test that export returns an ExportSummary."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter, ExportSummary

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",))
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "123"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        summary = exporter.export(str(filepath))

        assert summary is not None
        assert isinstance(summary, ExportSummary)
        assert summary.success == True

    def test_export_summary_sample_count(self, tmp_path):
        """Test that summary correctly reports sample count."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "1"}),
            create_test_sample("s2", "Sample2.fcs", {"Mouse #": "2"}),
            create_test_sample("s3", "Sample3.fcs", {"Mouse #": "3"}),
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        summary = exporter.export(str(filepath))

        assert summary.samples_processed == 3

    def test_export_summary_gate_count(self, tmp_path):
        """Test that summary correctly reports gate count."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([
            create_test_gate("cells", "Cells", ("Cells",)),
            create_test_gate("live", "Live", ("Cells", "Live")),
            create_test_gate("cd4", "CD4+", ("Cells", "Live", "CD4+")),
        ])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "1"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=["cells", "live", "cd4"],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        summary = exporter.export(str(filepath))

        assert summary.gates_exported == 3

    def test_export_summary_sheets_created(self, tmp_path):
        """Test that summary lists all created sheets."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "1"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=[],
            all_markers=[],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=True
        )

        filepath = tmp_path / "test.xlsx"
        summary = exporter.export(str(filepath))

        assert "Gating_Strategy" in summary.sheets_created
        assert "Keywords" in summary.sheets_created
        assert "Keywords_Wide" in summary.sheets_created

    def test_export_summary_mfi_metrics(self, tmp_path):
        """Test that summary correctly reports MFI metrics."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "1"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=["CD4", "CD8"],
            all_markers=["CD4", "CD8"],
            mfi_metrics=[MFIMetric.MEAN, MFIMetric.MEDIAN, MFIMetric.GEOMETRIC_MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        summary = exporter.export(str(filepath))

        assert len(summary.statistics.mfi_metrics) == 3
        assert "Mean" in summary.statistics.mfi_metrics
        assert "Median" in summary.statistics.mfi_metrics
        assert "GeoMean" in summary.statistics.mfi_metrics

    def test_export_summary_channel_count(self, tmp_path):
        """Test that summary correctly reports channel count."""
        from flowkit_gui.exporters.xlsx_exporter import XLSXExporter

        workspace = MockWorkspaceModel()
        workspace.set_gates([])
        workspace.set_samples([
            create_test_sample("s1", "Sample1.fcs", {"Mouse #": "1"})
        ])

        keywords_model = KeywordsModel()
        keywords_model.set_samples(workspace.samples)

        exporter = XLSXExporter(
            workspace=workspace,
            keywords_model=keywords_model,
            gate_ids=[],
            selected_markers=["CD4", "CD8"],
            all_markers=["CD4", "CD8", "CD45", "CD3"],
            mfi_metrics=[MFIMetric.MEAN],
            include_wide_keywords=False
        )

        filepath = tmp_path / "test.xlsx"
        summary = exporter.export(str(filepath))

        # Should include unique channels from both selected and all markers
        assert summary.statistics.channels_analyzed == 4


class TestXLSXExporterValueSanitization:
    """Tests for _sanitize_value ensuring safe openpyxl serialization."""

    def setup_method(self):
        import numpy as np
        self.np = np

    def test_nan_returns_empty_string(self):
        assert XLSXExporter._sanitize_value(float('nan')) == ""

    def test_none_returns_empty_string(self):
        assert XLSXExporter._sanitize_value(None) == ""

    def test_np_nan_returns_empty_string(self):
        assert XLSXExporter._sanitize_value(self.np.nan) == ""

    def test_inf_returns_empty_string(self):
        assert XLSXExporter._sanitize_value(float('inf')) == ""

    def test_neg_inf_returns_empty_string(self):
        assert XLSXExporter._sanitize_value(float('-inf')) == ""

    def test_np_inf_returns_empty_string(self):
        assert XLSXExporter._sanitize_value(self.np.inf) == ""

    def test_np_int64_returns_python_int(self):
        val = XLSXExporter._sanitize_value(self.np.int64(42))
        assert val == 42
        assert type(val) is int

    def test_np_float64_returns_python_float(self):
        val = XLSXExporter._sanitize_value(self.np.float64(3.14))
        assert val == 3.14
        assert type(val) is float

    def test_normal_int_passes_through(self):
        assert XLSXExporter._sanitize_value(42) == 42

    def test_normal_float_passes_through(self):
        assert XLSXExporter._sanitize_value(3.14) == 3.14

    def test_normal_string_passes_through(self):
        assert XLSXExporter._sanitize_value("hello") == "hello"

    def test_long_string_truncated(self):
        long_str = "x" * 40000
        result = XLSXExporter._sanitize_value(long_str)
        assert len(result) == 32767


class TestSanitizeSheetName:
    """Tests for _sanitize_sheet_name ensuring valid Excel sheet titles."""

    def test_plain_name_unchanged(self):
        assert XLSXExporter._sanitize_sheet_name("Freq_of_CD45") == "Freq_of_CD45"

    def test_colon_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("Freq_of_CD45+: Activated")
        assert ":" not in result
        assert result == "Freq_of_CD45+_ Activated"

    def test_slash_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("Bead_Norm (1000 beads/well)")
        assert "/" not in result
        assert result == "Bead_Norm (1000 beads_well)"

    def test_backslash_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("Gate\\Sub")
        assert "\\" not in result
        assert result == "Gate_Sub"

    def test_question_mark_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("Gate?Name")
        assert "?" not in result
        assert result == "Gate_Name"

    def test_asterisk_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("Gate*Name")
        assert "*" not in result
        assert result == "Gate_Name"

    def test_brackets_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("Gate[0]")
        assert "[" not in result
        assert "]" not in result
        assert result == "Gate_0_"

    def test_all_forbidden_chars_replaced(self):
        result = XLSXExporter._sanitize_sheet_name("a:b\\c/d?e*f[g]")
        assert result == "a_b_c_d_e_f_g_"

    def test_truncation_to_31_chars(self):
        long_name = "A" * 50
        result = XLSXExporter._sanitize_sheet_name(long_name)
        assert len(result) == 31

    def test_leading_apostrophe_stripped(self):
        result = XLSXExporter._sanitize_sheet_name("'LeadingQuote")
        assert not result.startswith("'")
        assert result == "LeadingQuote"

    def test_whitespace_stripped(self):
        result = XLSXExporter._sanitize_sheet_name("  Spaced  ")
        assert result == "Spaced"

    def test_empty_string(self):
        result = XLSXExporter._sanitize_sheet_name("")
        assert result == ""

    def test_realistic_gate_with_colon(self):
        """Simulates the exact crash scenario Colleen reported."""
        gate_name = "CD45+: Activated"
        result = XLSXExporter._sanitize_sheet_name(f"Freq_of_{gate_name}")
        assert len(result) <= 31
        assert ":" not in result
        assert result == "Freq_of_CD45+_ Activated"

    def test_realistic_bead_per_well(self):
        """Catches the latent bug in bead 'per_well' mode."""
        result = XLSXExporter._sanitize_sheet_name("Bead_Norm (50000 beads/well)")
        assert "/" not in result
        assert len(result) <= 31


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
